# -*- coding: utf-8 -*-
"""
Tests GPV-F-77 — Formato Estudio de Títulos Art. 276 Ley 1955 de 2019.

Verifican que el DOCX oficial del Ministerio de Vivienda se pre-diligencia con
los datos reales del CTL (identificación, tracto sucesivo cronológico,
titularidad, limitaciones) y que, cuando un dato no tiene fuente, se marca
"NO REGISTRA" / "NO APLICA (N/A)" — regla del instructivo GPV-I-20 — sin
inventar información. La plantilla oficial nunca se modifica en disco.
"""
import sys
import unittest
from io import BytesIO
from pathlib import Path

import pytest

ROOT_DIR = Path(__file__).resolve().parent.parent
API_DIR = ROOT_DIR / "api"
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(API_DIR))

_PLANTILLA = API_DIR / "plantillas" / "GPV-F-77_Estudio_Titulos_Art276_Ley1955_2019.docx"

# Texto mínimo de un CTL SNR real con los campos que hoy importan
_TEXTO_CTL_REAL = """
CERTIFICADO DE TRADICION Y LIBERTAD
MATRICULA INMOBILIARIA: 040-347004
CIRCULO REGISTRAL: 040 - BARRANQUILLA
FECHA APERTURA: 28-03-2001
CODIGO CATASTRAL: 080010102000001710004000000000
NUPRE: AFT0040BBHC
ESTADO DEL FOLIO: ACTIVO
DESCRIPCION: CABIDA Y LINDEROS
BODEGA NUMERO 3, CON AREA DE 570.9 MTS2.
DIRECCION DEL INMUEBLE
Tipo Predio: URBANO
2) CL 40 # 50 B - 62 BDGA 3
ANOTACION: Nro 001 Fecha: 21-03-2001
ESPECIFICACION: MODO DE ADQUISICION: 160 DIVISION MATERIAL
ESCRITURA PUBLICA No. 1234 DEL 15-03-2001 NOTARIA QUINTA DE BARRANQUILLA
DE: URBANIZADORA EJEMPLO S.A.
A: INDUSTRIAS DEL NORTE LTDA.
ANOTACION: Nro 002 Fecha: 03-03-2009
ESPECIFICACION: GRAVAMEN: 0205 HIPOTECA
A: BANCO EJEMPLO S.A.
ANOTACION: Nro 006 Fecha: 09-08-2017
Se cancela anotación No: 2
ESPECIFICACION: CANCELACION: 0843 CANCELACION HIPOTECA
NRO TOTAL DE ANOTACIONES: *6*
"""


def _analizar_y_generar(extra=None, sin_ctl=False):
    """Analiza el CTL de prueba y genera el DOCX. Devuelve (ctx, doc)."""
    from legal_analyzer import analizar_texto_certificado
    from gpv_f77 import generar_docx_gpv_f77
    from docx import Document

    if sin_ctl:
        analysis = analizar_texto_certificado(None)
        ctx = {
            "folio": "040-347004", "ciudad": "barranquilla", "sin_ctl": True,
            "direccion": "", "barrio": "", "area_juridica": None, "area_catastral": None,
            "codigo_catastral": None, "nupre": None, "apertura": "",
            "titulares": "", "descripcion_ctl": "", "anotaciones_detalle": [],
        }
    else:
        analysis = analizar_texto_certificado(_TEXTO_CTL_REAL)
        ctx = {
            "folio": analysis["folio"],
            "ciudad": "barranquilla",
            "direccion": "CL 40 # 50 B - 62 BDGA 3",
            "barrio": "Barrio Abajo",
            "area_juridica": 570.9,
            "area_catastral": 475.0,
            "codigo_catastral": analysis["codigo_catastral"],
            "nupre": analysis["nupre"],
            "apertura": analysis["apertura"],
            "titulares": analysis["titulares"],
            "descripcion_ctl": analysis["descripcion_ctl"],
            "anotaciones_detalle": analysis.get("anotaciones_detalle") or [],
            "sin_ctl": False,
        }
    if extra:
        ctx.update(extra)
    data = generar_docx_gpv_f77(ctx)
    doc = Document(BytesIO(data))
    return ctx, doc, data


class TestGpvF77Identificacion(unittest.TestCase):

    def test_plantilla_existe_en_api(self):
        self.assertTrue(_PLANTILLA.exists(),
                        "plantilla GPV-F-77 debe empaquetarse con api/ (Vercel)")

    def test_identificacion_con_datos_del_ctl(self):
        """La Sección 1 se llena con folio, departamento/municipio, dirección,
        barrio, matrícula, código catastral y áreas reales del CTL."""
        ctx, doc, _ = _analizar_y_generar()
        t0 = doc.tables[0]
        self.assertTrue(t0.rows[1].cells[1].text.strip().startswith("N/A"),
                        "expediente interno ICT-INURBE no suministrado -> N/A")
        self.assertRegex(t0.rows[1].cells[3].text.strip(),
                         r"^\d{2}-[a-z]{3}-\d{4}$")  # dd-mmm-aaaa
        self.assertEqual(t0.rows[2].cells[1].text.strip(), "Atlántico")
        self.assertEqual(t0.rows[2].cells[3].text.strip(), "Barranquilla")
        self.assertIn("40", t0.rows[3].cells[1].text)
        self.assertIn("Barrio Abajo", t0.rows[3].cells[3].text)
        self.assertEqual(t0.rows[4].cells[1].text.strip(), "040-347004")
        self.assertIn("N/A", t0.rows[4].cells[3].text)  # sin segregación
        self.assertEqual(t0.rows[5].cells[1].text.strip(),
                         "080010102000001710004000000000")
        self.assertIn("570.9", t0.rows[6].cells[1].text)
        self.assertIn("475.0", t0.rows[6].cells[1].text)

    def test_ciudad_medellin_escribe_antioquia(self):
        ctx, doc, _ = _analizar_y_generar(extra={"ciudad": "medellin"})
        t0 = doc.tables[0]
        self.assertEqual(t0.rows[2].cells[1].text.strip(), "Antioquia")
        self.assertEqual(t0.rows[2].cells[3].text.strip(), "Medellín")

    def test_ciudad_bogota_escribe_distrito(self):
        ctx, doc, _ = _analizar_y_generar(extra={"ciudad": "bogota"})
        t0 = doc.tables[0]
        self.assertEqual(t0.rows[2].cells[1].text.strip(), "Bogotá D.C.")
        self.assertEqual(t0.rows[2].cells[3].text.strip(), "Bogotá D.C.")


class TestGpvF77Juridico(unittest.TestCase):

    def test_tracto_sucesivo_cronologico_con_escritura(self):
        """El tracto menciona la apertura y la anotación de dominio con su
        escritura/notaría y las partes, en orden cronológico."""
        ctx, doc, _ = _analizar_y_generar()
        t1 = doc.tables[1]
        texto_celda = "\n".join(p.text for p in t1.rows[2].cells[0].paragraphs)
        self.assertIn("Folio abierto el", texto_celda)
        self.assertIn("Anotación Nro. 001 (21-03-2001)", texto_celda)
        self.assertIn("DIVISION MATERIAL", texto_celda.upper())
        self.assertIn("Escritura No. 1234", texto_celda)
        self.assertIn("a favor de: INDUSTRIAS DEL NORTE LTDA", texto_celda)

    def test_linderos_desde_descripcion_ctl(self):
        ctx, doc, _ = _analizar_y_generar()
        t1 = doc.tables[1]
        texto_celda = "\n".join(p.text for p in t1.rows[2].cells[0].paragraphs)
        self.assertIn("BODEGA NUMERO 3", texto_celda.upper())

    def test_titularidad_actual(self):
        """La titularidad actual proviene del CTL, no se inventa."""
        ctx, doc, _ = _analizar_y_generar()
        t1 = doc.tables[1]
        self.assertEqual(t1.rows[4].cells[0].text.strip(),
                         "INDUSTRIAS DEL NORTE LTDA")

    def test_hipoteca_cancelada_no_es_gravamen_vigente(self):
        """La hipoteca anot. 002 fue cancelada (anot. 006): el estudio NO la
        declara vigente; aparece en el historial de canceladas (regresión)."""
        ctx, doc, _ = _analizar_y_generar()
        t1 = doc.tables[1]
        texto = t1.rows[8].cells[0].text
        self.assertIn("NO REGISTRA gravámenes", texto)
        self.assertIn("Historial de limitaciones canceladas", texto)
        self.assertIn("Hipoteca", texto)
        self.assertIn("cancelada por la anotación Nro. 006", texto)
        # La recomendación no debe bloquear por el gravamen ya cancelado
        rec = t1.rows[12].cells[1].text
        self.assertIn("POSITIVA", rec)
        self.assertNotIn("VIGENTES", rec)

    def test_impuestos_sin_fuente_marcan_no_registra(self):
        """Sin paz y salvo (fuente que allega la entidad pública) -> NO REGISTRA."""
        ctx, doc, _ = _analizar_y_generar()
        t1 = doc.tables[1]
        self.assertIn("NO REGISTRA", t1.rows[6].cells[0].text)

    def test_documentos_listados_con_folio_real(self):
        ctx, doc, _ = _analizar_y_generar()
        t1 = doc.tables[1]
        texto_docs = "\n".join(p.text for p in t1.rows[11].cells[1].paragraphs)
        self.assertIn("Consulta VUR FMI 040-347004", texto_docs)
        self.assertIn("Escritura No. 1234", texto_docs)

    def test_recomendacion_borrador_con_advertencia_sin_ctl(self):
        ctx, doc, _ = _analizar_y_generar(sin_ctl=True)
        t1 = doc.tables[1]
        self.assertIn("NO APTO", t1.rows[12].cells[1].text)


class TestGpvF77Honestidad(unittest.TestCase):

    def test_sin_ctl_todo_es_no_registra(self):
        """Sin CTL no se afirma nada: cada campo queda NO REGISTRA / N/A."""
        ctx, doc, _ = _analizar_y_generar(sin_ctl=True)
        t0 = doc.tables[0]
        self.assertIn("NO REGISTRA", t0.rows[3].cells[1].text)
        self.assertIn("NO REGISTRA", t0.rows[4].cells[1].text)
        self.assertIn("NO REGISTRA", t0.rows[5].cells[1].text)
        self.assertIn("NO REGISTRA", t0.rows[6].cells[1].text)
        t1 = doc.tables[1]
        self.assertIn("NO REGISTRA", t1.rows[4].cells[0].text)
        self.assertIn("NO REGISTRA", t1.rows[6].cells[0].text)
        self.assertIn("NO REGISTRA", t1.rows[8].cells[0].text)

    def test_plantilla_oficial_no_se_modifica_en_disco(self):
        """El generador trabaja en memoria: el archivo de plantilla del repo
        queda intacto byte a byte (documento controlado, copia nunca alterada)."""
        antes = _PLANTILLA.read_bytes()
        _analizar_y_generar()
        _analizar_y_generar(sin_ctl=True)
        despues = _PLANTILLA.read_bytes()
        self.assertEqual(antes, despues)

    def test_seccion_4_queda_para_firma_manual(self):
        """Los datos del profesional que diligencia quedan en blanco: la firma
        (Sección 4) es responsabilidad del profesional jurídico, no automática."""
        ctx, doc, _ = _analizar_y_generar()
        t2 = doc.tables[2]
        self.assertEqual(t2.rows[1].cells[1].text.strip(), "")   # Diligenciado por
        self.assertEqual(t2.rows[2].cells[1].text.strip(), "")   # Profesión
        self.assertEqual(t2.rows[2].cells[3].text.strip(), "")   # Tarjeta profesional


class TestGpvF77Endpoint(unittest.TestCase):

    def test_endpoint_requiere_certificado(self):
        """Sin CTL el endpoint responde 400: el Estudio de Títulos no tiene
        fuente registral (regla GPV-I-20: verificación previa de documentos)."""
        from fastapi import HTTPException
        from index import generar_gpv_f77_endpoint
        import asyncio

        async def _correr():
            with self.assertRaises(HTTPException) as ctx:
                await generar_gpv_f77_endpoint(
                    folio_matricula="040-347004", direccion="CL 40 # 50 B - 62",
                    area=570.9, barrio="Barrio Abajo", ciudad="barranquilla",
                    certificado=None, auth={"username": "admin", "rol": "admin"})
            self.assertEqual(ctx.exception.status_code, 400)

        asyncio.run(_correr())

    @pytest.mark.network
    def test_endpoint_genera_docx_con_certificado(self):
        """En vivo (marcado network): con un CTL mínimo el endpoint devuelve un
        DOCX con la identificación llena (geocodifica contra el servicio público)."""
        import asyncio
        import os
        import tempfile
        from pathlib import Path as _P

        # Dir temporal dentro del repo (los tests no escriben en /tmp)
        tmp = ROOT_DIR / "tmp_gpv77_test"
        tmp.mkdir(parents=True, exist_ok=True)
        os.environ["ARHIAX_TMP_DIR"] = str(tmp)

        from reportlab.pdfgen import canvas
        buf = BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(50, 780, "CERTIFICADO DE TRADICION Y LIBERTAD (PRUEBA)")
        c.drawString(50, 760, "MATRICULA INMOBILIARIA: 040-777777")
        c.drawString(50, 740, "DIRECCION: CALLE 63 # 37-71")
        c.drawString(50, 720, "AREA PRIVADA 75.5 M2")
        c.save()
        class FakeFile:
            def __init__(self, data):
                self._data = data
                self.filename = "cert.pdf"
                self.file = BytesIO(data)
        cert = FakeFile(buf.getvalue())

        from index import generar_gpv_f77_endpoint
        async def _correr():
            return await generar_gpv_f77_endpoint(
                folio_matricula=None, direccion=None, area=None, barrio=None,
                ciudad="barranquilla", certificado=cert,
                auth={"username": "admin", "rol": "admin"})
        try:
            resp = asyncio.run(_correr())
            self.assertIn("wordprocessingml", resp.media_type)
            self.assertIn("GPV-F-77_Estudio_Titulos", resp.headers["Content-Disposition"])
            from docx import Document
            doc = Document(BytesIO(resp.body))
            t0 = doc.tables[0]
            self.assertIn("040-777777", t0.rows[4].cells[1].text)
        finally:
            import shutil
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
